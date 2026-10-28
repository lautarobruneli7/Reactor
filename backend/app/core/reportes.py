"""Reportes de ensayos en PDF y Excel.

El PDF usa la tipografía estándar (Helvetica), que no tiene subíndices ni flechas: `_t` los reemplaza por texto plano
(H₂O → H2O, → → ->, ≥ → >=). Un reporte nunca dice más de lo que dice el ensayo: el origen de cada dato
(estimado / medido), el estado de validación del modelo y la huella de trazabilidad van siempre.
"""
from __future__ import annotations

import io
from datetime import datetime
from xml.sax.saxutils import escape as _xml

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app import config
from app.models.ensayos import Comparacion, EnsayoOut
from app.models.schemas import Compuesto

_SUB = str.maketrans("₀₁₂₃₄₅₆₇₈₉", "0123456789")
_REEMPLAZOS = {"→": "->", "≥": ">=", "≤": "<=", "·": ".", "–": "-", "—": "-", "‑": "-", "µ": "u"}
ORIGEN = {"estimacion": "ESTIMACION (modelo sin validar)", "mixto": "MIXTO (datos medidos y estimados)", "medido": "MEDIDO en laboratorio"}
ESTADO_ESPEC = {"cumple": "Cumple", "no_cumple": "NO cumple", "sin_dato": "Sin dato", "sin_limite": "Sin limite cargado"}
GLOBAL_ESPEC = {
    "cumple": "Cumple la especificacion en todos los parametros.",
    "no_cumple": "NO cumple la especificacion.",
    "incompleto": "Especificacion INCOMPLETA: faltan mediciones o limites; no se puede afirmar que el producto cumpla.",
    "sin_especificaciones": "El producto no tiene una especificacion cargada.",
}
AVISO_FINAL = (
    "Herramienta de estimacion con datos simplificados; no reemplaza ensayos de laboratorio. El puntaje mide la calidad de la mezcla "
    "(pureza, rendimiento, balance estequiometrico), no la eficacia clinica. Los valores deben validarse con fuentes oficiales "
    "(USP, Ph. Eur., Farmacopea Argentina) y un profesional del area."
)


def _t(texto: object) -> str:
    """Texto seguro para un párrafo del PDF: tipografía estándar (Latin-1) y sin marcas (&, <, >) que el PDF tomaría por formato."""
    s = str(texto).translate(_SUB)
    for a, b in _REEMPLAZOS.items():
        s = s.replace(a, b)
    return _xml(s.encode("latin-1", "replace").decode("latin-1"))


def _num(v: float | None, dec: int = 2) -> str:
    return "-" if v is None else f"{v:.{dec}f}"


def _cantidad(c: dict) -> str:
    if c.get("gramos") is not None:
        return f"{c['gramos']:g} g"
    return f"{c['cantidad']:g} {c['unidad']}"


def _concentracion(c: dict) -> str:
    k = c.get("concentracion")
    if not k:
        return "-"
    unidad = {"p/p": "% p/p", "p/v": "% p/v", "molar": "mol/L"}[k["tipo"]]
    return f"{k['valor']:g} {unidad}"


def _fecha(iso: str) -> str:
    try:
        return datetime.fromisoformat(iso).strftime("%d/%m/%Y %H:%M UTC")
    except ValueError:
        return iso


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def reporte_pdf(e: EnsayoOut, catalogo: dict[str, Compuesto]) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
        title=f"Retorta - Ensayo {e.id:03d}", author="Retorta",
    )
    base = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=base["Heading1"], fontSize=17, spaceAfter=4)
    h2 = ParagraphStyle("h2", parent=base["Heading2"], fontSize=12, spaceBefore=10, spaceAfter=4, textColor=colors.HexColor("#0b3b66"))
    p = ParagraphStyle("p", parent=base["BodyText"], fontSize=9, leading=12)
    chico = ParagraphStyle("chico", parent=p, fontSize=7.5, leading=10, textColor=colors.HexColor("#444444"))
    aviso = ParagraphStyle("aviso", parent=p, backColor=colors.HexColor("#fff4de"), borderColor=colors.HexColor("#e0a23a"), borderWidth=0.6, borderPadding=5, spaceBefore=4, spaceAfter=6)

    def tabla(filas: list[list], anchos: list[float] | None = None, cabecera: bool = True) -> Table:
        t = Table([[Paragraph(_t(c), p) for c in fila] for fila in filas], colWidths=anchos, repeatRows=1 if cabecera else 0)
        est = [("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#b8c4d6")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("FONTSIZE", (0, 0), (-1, -1), 8.5)]
        if cabecera:
            est += [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e6eefc"))]
        t.setStyle(TableStyle(est))
        return t

    r = e.salida.resultado
    assert r is not None
    out: list = []
    out.append(Paragraph(_t(f"Retorta - Reporte del ensayo N° {e.id:03d}"), h1))
    out.append(Paragraph(_t(f"{r.reaccion_nombre}"), p))
    out.append(Spacer(1, 4))
    out.append(tabla([
        ["Fecha", _fecha(e.creado_en), "Realizado por", e.usuario],
        ["Lote", e.lote or "-", "Estado", e.estado + ("" if e.vigente else f" (reemplazado por N° {e.reemplazado_por})")],
        ["Version del modelo", e.version_modelo, "Corrige al ensayo", f"N° {e.corrige_a}" if e.corrige_a else "-"],
    ], [32 * mm, 56 * mm, 32 * mm, 52 * mm], cabecera=False))

    out.append(Paragraph("Origen de los datos", h2))
    out.append(Paragraph("<b>" + _t(ORIGEN[r.origen]) + "</b>", p))
    if r.origen != "medido":
        out.append(Paragraph(_t("Los valores marcados como estimados salen de un modelo cuyos parametros no estan validados. " + (r.estado_validacion or "")), aviso))

    out.append(Paragraph("Carga", h2))
    filas = [["Compuesto", "Cantidad", "Pureza", "Concentracion", "Densidad", "Material (g)", "Moles netos"]]
    cargas = {c.id: c for c in r.cargas}
    for c in e.entrada.componentes:
        carga = cargas.get(c.id)
        d = c.model_dump()
        filas.append([
            catalogo[c.id].nombre if c.id in catalogo else c.id, _cantidad(d), f"{c.pureza:g} %", _concentracion(d),
            f"{c.densidad_g_ml:g} g/mL" if c.densidad_g_ml else "-",
            _num(carga.gramos_material, 3) if carga else "no participa",
            _num(carga.moles_netos, 4) if carga and carga.moles_netos is not None else "-",
        ])
    out.append(tabla(filas, [38 * mm, 24 * mm, 16 * mm, 27 * mm, 22 * mm, 25 * mm, 22 * mm]))

    out.append(Paragraph("Resultados", h2))
    tag = lambda o: "medido" if o == "experimental" else "estimado"  # noqa: E731
    out.append(tabla([
        ["Puntaje de calidad", f"{r.puntaje}/100 ({r.categoria})"],
        ["Rendimiento", f"{_num(r.rendimiento_pct, 1)} %  [{tag(r.rendimiento_origen)}]"],
        ["Pureza del producto", f"{_num(r.pureza_producto_pct, 1)} %  [{tag(r.pureza_origen)}]"],
        ["Balance estequiometrico", f"{_num(r.balance_pct, 0)} / 100"],
        ["Producto", f"{_num(r.gramos_estimados, 2)} g de {_num(r.gramos_teoricos, 2)} g teoricos"],
        ["Reactivo limitante", r.reactivo_limitante],
        ["Economia atomica", f"{_num(r.economia_atomica_pct, 1)} %"],
        ["Ecuacion", r.ecuacion],
    ], [48 * mm, 126 * mm], cabecera=False))

    if r.incertidumbre:
        inc = r.incertidumbre
        out.append(Paragraph("Incertidumbre (Monte Carlo)", h2))
        out.append(tabla(
            [["Magnitud", "Media", "p5", "p50", "p95"]]
            + [[n, str(iv.media), str(iv.p5), str(iv.p50), str(iv.p95)] for n, iv in (
                ("Rendimiento (%)", inc.rendimiento_pct), ("Pureza (%)", inc.pureza_producto_pct),
                ("Producto (g)", inc.gramos_estimados), ("Puntaje", inc.puntaje))],
            [54 * mm, 30 * mm, 30 * mm, 30 * mm, 30 * mm],
        ))
        out.append(Paragraph(_t(f"{inc.corridas_validas} corridas, semilla {inc.semilla}. Probabilidad de cumplir el minimo de pureza: {inc.prob_cumple_pureza * 100:.1f} %."), p))
        for n in inc.notas:
            out.append(Paragraph(_t("- " + n), chico))

    out.append(Paragraph("Especificacion del producto", h2))
    out.append(Paragraph(_t(GLOBAL_ESPEC[r.cumplimiento_especificaciones]), p))
    if r.especificaciones:
        def lim(s) -> str:
            u = f" {s.unidad}" if s.unidad else ""
            if s.minimo is not None and s.maximo is not None:
                return f"{s.minimo:g} - {s.maximo:g}{u}"
            if s.minimo is not None:
                return f">= {s.minimo:g}{u}"
            if s.maximo is not None:
                return f"<= {s.maximo:g}{u}"
            return "sin limite"
        out.append(tabla(
            [["Parametro", "Limite", "Valor", "Origen", "Estado"]]
            + [[s.parametro, lim(s), _num(s.valor, 2), s.origen or "-", ESTADO_ESPEC[s.estado]] for s in r.especificaciones],
            [62 * mm, 30 * mm, 24 * mm, 22 * mm, 36 * mm],
        ))

    if e.salida.alertas:
        out.append(Paragraph("Alertas", h2))
        for a in e.salida.alertas:
            out.append(Paragraph(_t(f"[{a.tipo.upper()}] {a.mensaje}"), p))

    if e.notas:
        out.append(Paragraph("Notas", h2))
        out.append(Paragraph(_t(e.notas).replace("\n", "<br/>"), p))

    if e.eventos:
        out.append(Paragraph("Revision", h2))
        out.append(tabla(
            [["Fecha", "Evento", "Persona", "Comentario"]] + [[_fecha(v.creado_en), v.tipo, v.usuario, v.comentario or "-"] for v in e.eventos],
            [36 * mm, 24 * mm, 36 * mm, 78 * mm],
        ))

    out.append(Paragraph("Trazabilidad", h2))
    traza = [["Huella del ensayo (SHA-256)", e.digest]]
    if e.sello:
        traza += [[f"Registro N° {e.sello.id}: hash", e.sello.hash], ["Hash del registro anterior", e.sello.prev_hash]]
    out.append(tabla(traza, [50 * mm, 124 * mm], cabecera=False))
    out.append(Paragraph(_t(f"Generado por Retorta {config.VERSION}. Verificable con GET /api/auditoria/verificar."), chico))
    out.append(Spacer(1, 6))
    out.append(Paragraph(_t(AVISO_FINAL), chico))

    doc.build(out)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
_NEGRITA = Font(bold=True)
_ENCABEZADO = PatternFill("solid", fgColor="E6EEFC")


def _encabezado(ws, fila: int, valores: list[str]) -> None:
    for j, v in enumerate(valores, start=1):
        c = ws.cell(row=fila, column=j, value=v)
        c.font = _NEGRITA
        c.fill = _ENCABEZADO
        c.alignment = Alignment(wrap_text=True, vertical="top")


def _ajustar(ws, anchos: dict[int, int] | None = None) -> None:
    for col in ws.columns:
        letra = get_column_letter(col[0].column)
        ancho = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[letra].width = min(60, max(10, ancho + 2)) if not anchos or col[0].column not in anchos else anchos[col[0].column]


def _hojas_ensayo(wb: Workbook, e: EnsayoOut, catalogo: dict[str, Compuesto], primera: bool) -> None:
    r = e.salida.resultado
    assert r is not None
    ws = wb.active if primera else wb.create_sheet()
    ws.title = f"Ensayo {e.id:03d}"[:31]
    filas = [
        ("Ensayo N°", e.id), ("Fecha (UTC)", e.creado_en), ("Realizado por", e.usuario), ("Lote", e.lote or ""),
        ("Reacción", r.reaccion_nombre), ("Ecuación", r.ecuacion), ("Estado", e.estado), ("Vigente", "sí" if e.vigente else f"no: reemplazado por N° {e.reemplazado_por}"),
        ("Versión del modelo", e.version_modelo), ("Origen de los datos", ORIGEN[r.origen]), ("Estado de validación del modelo", r.estado_validacion),
        ("Puntaje (0-100)", r.puntaje), ("Categoría", r.categoria),
        ("Rendimiento (%)", r.rendimiento_pct), ("Origen del rendimiento", "medido" if r.rendimiento_origen == "experimental" else "estimado por el modelo"),
        ("Pureza del producto (%)", r.pureza_producto_pct), ("Origen de la pureza", "medida" if r.pureza_origen == "experimental" else "estimada por el modelo"),
        ("Balance (0-100)", r.balance_pct), ("Producto estimado (g)", r.gramos_estimados), ("Producto teórico (g)", r.gramos_teoricos),
        ("Reactivo limitante", r.reactivo_limitante), ("Economía atómica (%)", r.economia_atomica_pct),
        ("Especificación", GLOBAL_ESPEC[r.cumplimiento_especificaciones]), ("Notas", e.notas),
    ]
    _encabezado(ws, 1, ["Dato", "Valor"])
    for i, (k, v) in enumerate(filas, start=2):
        ws.cell(row=i, column=1, value=k).font = _NEGRITA
        ws.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    _ajustar(ws, {1: 34, 2: 70})

    ws = wb.create_sheet(f"Carga {e.id:03d}"[:31])
    _encabezado(ws, 1, ["Compuesto", "Cantidad", "Unidad", "Pureza (%)", "Concentración", "Densidad (g/mL)", "Material (g)", "Fracción de soluto", "Moles netos"])
    cargas = {c.id: c for c in r.cargas}
    for i, c in enumerate(e.entrada.componentes, start=2):
        carga = cargas.get(c.id)
        d = c.model_dump()
        cant = d["gramos"] if d.get("gramos") is not None else d["cantidad"]
        uni = "g" if d.get("gramos") is not None else d["unidad"]
        for j, v in enumerate([
            catalogo[c.id].nombre if c.id in catalogo else c.id, cant, uni, c.pureza, _concentracion(d), c.densidad_g_ml,
            carga.gramos_material if carga else None, carga.fraccion_soluto if carga else None, carga.moles_netos if carga else None,
        ], start=1):
            ws.cell(row=i, column=j, value=v)
    _ajustar(ws)

    ws = wb.create_sheet(f"Especificación {e.id:03d}"[:31])
    _encabezado(ws, 1, ["Parámetro", "Unidad", "Mínimo", "Máximo", "Valor", "Origen", "Estado", "Fuente / nota"])
    for i, s in enumerate(r.especificaciones, start=2):
        for j, v in enumerate([s.parametro, s.unidad, s.minimo, s.maximo, s.valor, s.origen, ESTADO_ESPEC[s.estado], s.fuente or s.estado_limite], start=1):
            ws.cell(row=i, column=j, value=v)
    _ajustar(ws)

    if r.incertidumbre:
        inc = r.incertidumbre
        ws = wb.create_sheet(f"Incertidumbre {e.id:03d}"[:31])
        _encabezado(ws, 1, ["Magnitud", "Media", "Desvío", "p5", "p50", "p95"])
        for i, (n, iv) in enumerate((
            ("Rendimiento (%)", inc.rendimiento_pct), ("Pureza (%)", inc.pureza_producto_pct),
            ("Producto (g)", inc.gramos_estimados), ("Puntaje", inc.puntaje)), start=2):
            for j, v in enumerate([n, iv.media, iv.sd, iv.p5, iv.p50, iv.p95], start=1):
                ws.cell(row=i, column=j, value=v)
        base = 8
        ws.cell(row=base, column=1, value="Corridas válidas").font = _NEGRITA
        ws.cell(row=base, column=2, value=inc.corridas_validas)
        ws.cell(row=base + 1, column=1, value="Semilla").font = _NEGRITA
        ws.cell(row=base + 1, column=2, value=inc.semilla)
        ws.cell(row=base + 2, column=1, value="Probabilidad de cumplir el mínimo de pureza").font = _NEGRITA
        ws.cell(row=base + 2, column=2, value=inc.prob_cumple_pureza)
        for k, n in enumerate(inc.notas):
            ws.cell(row=base + 4 + k, column=1, value=n)
        _ajustar(ws, {1: 44})

    ws = wb.create_sheet(f"Trazabilidad {e.id:03d}"[:31])
    _encabezado(ws, 1, ["Dato", "Valor"])
    tz = [("Huella del ensayo (SHA-256)", e.digest)]
    if e.sello:
        tz += [("N° de registro", e.sello.id), ("Hash del registro", e.sello.hash), ("Hash del registro anterior", e.sello.prev_hash), ("Registrado (UTC)", e.sello.creado_en)]
    for v in e.eventos:
        tz.append((f"{v.tipo} — {v.usuario} ({v.creado_en})", v.comentario))
    for i, (k, v) in enumerate(tz, start=2):
        ws.cell(row=i, column=1, value=k).font = _NEGRITA
        ws.cell(row=i, column=2, value=v)
    _ajustar(ws, {1: 50, 2: 70})


def reporte_xlsx(e: EnsayoOut, catalogo: dict[str, Compuesto]) -> bytes:
    wb = Workbook()
    _hojas_ensayo(wb, e, catalogo, primera=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def reporte_comparacion_xlsx(c: Comparacion, catalogo: dict[str, Compuesto]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparación"
    ids = [e.id for e in c.ensayos]
    _encabezado(ws, 1, ["Métrica"] + [f"N° {i:03d}" for i in ids])
    fila = 2
    ws.cell(row=fila, column=1, value="Fecha").font = _NEGRITA
    for j, e in enumerate(c.ensayos, start=2):
        ws.cell(row=fila, column=j, value=e.creado_en)
    fila += 1
    ws.cell(row=fila, column=1, value="Lote").font = _NEGRITA
    for j, e in enumerate(c.ensayos, start=2):
        ws.cell(row=fila, column=j, value=e.lote or "")
    fila += 1
    for m in c.metricas:
        ws.cell(row=fila, column=1, value=m["metrica"]).font = _NEGRITA
        for j, i in enumerate(ids, start=2):
            ws.cell(row=fila, column=j, value=m["valores"].get(i))
        fila += 1
    fila += 1
    _encabezado(ws, fila, ["Carga (g de material)"] + [f"N° {i:03d}" for i in ids])
    for comp in c.compuestos:
        fila += 1
        ws.cell(row=fila, column=1, value=catalogo[comp["id"]].nombre if comp["id"] in catalogo else comp["id"]).font = _NEGRITA
        for j, i in enumerate(ids, start=2):
            ws.cell(row=fila, column=j, value=comp["gramos"].get(i))
    _ajustar(ws, {1: 34})
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
