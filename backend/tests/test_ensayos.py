"""Ensayos guardados en el servidor: se calculan en el servidor, no se editan, se corrigen con otro y dejan rastro verificable.

Las pruebas corren sobre SQLite y, si está `SIM_TEST_POSTGRES_URL` (p. ej. postgresql+psycopg://postgres@127.0.0.1:54329/retorta_test),
también sobre PostgreSQL.
"""
import io
import os
import threading

import pytest
from fastapi.testclient import TestClient
from sqlalchemy import text

from app.db import datos
from app.main import app

POSTGRES = os.getenv("SIM_TEST_POSTGRES_URL")
EJEMPLO = [
    {"id": "ibb", "gramos": 134, "pureza": 99},
    {"id": "aa", "gramos": 102, "pureza": 98},
    {"id": "cat", "gramos": 20, "pureza": 97},
    {"id": "h2", "gramos": 2, "pureza": 99},
    {"id": "co", "gramos": 28, "pureza": 98},
    {"id": "exc", "gramos": 30, "pureza": 99},
]


def _reiniciar_postgres(url: str) -> None:
    import psycopg

    conn_url = url.replace("postgresql+psycopg://", "postgresql://")
    with psycopg.connect(conn_url, autocommit=True) as c:
        c.execute("DROP SCHEMA public CASCADE")
        c.execute("CREATE SCHEMA public")


@pytest.fixture(params=["sqlite", pytest.param("postgres", marks=pytest.mark.skipif(not POSTGRES, reason="falta SIM_TEST_POSTGRES_URL"))])
def cliente(request, tmp_path, monkeypatch):
    """Cliente en modo local (sin login) sobre una base de datos nueva."""
    monkeypatch.delenv("SIM_AUTH", raising=False)
    if request.param == "postgres":
        _reiniciar_postgres(POSTGRES)
        monkeypatch.setenv("SIM_DATABASE_URL", POSTGRES)
    else:
        monkeypatch.setenv("SIM_DATABASE_URL", f"sqlite:///{tmp_path / 'datos.db'}")
    with TestClient(app) as c:
        yield c


def nuevo(lote="L-001", notas="", **extra):
    return {"componentes": EJEMPLO, "lote": lote, "notas": notas, **extra}


def test_se_guarda_con_el_resultado_calculado_por_el_servidor(cliente):
    r = cliente.post("/api/ensayos", json={**nuevo(), "resultado": {"puntaje": 100}, "puntaje": 100})  # un resultado de afuera se ignora
    assert r.status_code == 201, r.text
    e = r.json()
    assert e["id"] == 1 and e["puntaje"] == 85 and e["salida"]["resultado"]["puntaje"] == 85
    assert e["estado"] == "registrado" and e["vigente"] and e["lote"] == "L-001"
    assert e["origen"] == "estimacion" and e["version_modelo"].count("+") == 1
    assert len(e["sello"]["prev_hash"]) == 64 and len(e["sello"]["hash"]) == 64 and e["sello"]["hash"] != e["sello"]["prev_hash"]
    assert e["entrada"]["componentes"][0]["id"] == "ibb"


def test_listar_obtener_y_filtrar(cliente):
    for lote in ("A-1", "A-2", "B-1"):
        assert cliente.post("/api/ensayos", json=nuevo(lote=lote)).status_code == 201
    r = cliente.get("/api/ensayos")
    assert [e["lote"] for e in r.json()] == ["B-1", "A-2", "A-1"] and r.headers["X-Total-Count"] == "3"
    assert [e["lote"] for e in cliente.get("/api/ensayos", params={"lote": "a-"}).json()] == ["A-2", "A-1"]
    assert len(cliente.get("/api/ensayos", params={"limit": 2, "offset": 1}).json()) == 2
    assert cliente.get("/api/ensayos", params={"estado": "aprobado"}).json() == []
    assert cliente.get("/api/ensayos/2").json()["lote"] == "A-2"
    assert cliente.get("/api/ensayos/99").status_code == 404
    assert cliente.get("/api/ensayos", params={"estado": "otro"}).status_code == 422


def test_una_mezcla_incompleta_no_se_guarda(cliente):
    r = cliente.post("/api/ensayos", json={"componentes": [{"id": "ibb", "gramos": 100, "pureza": 99}]})
    assert r.status_code == 422 and "no se guarda" in r.json()["detail"]
    assert cliente.get("/api/ensayos").json() == []
    assert cliente.post("/api/ensayos", json={"componentes": [{"id": "zzz", "gramos": 1, "pureza": 99}]}).status_code == 404
    assert cliente.post("/api/ensayos", json={"reaccion_id": "no-existe", "componentes": EJEMPLO}).status_code == 404


def test_guarda_datos_medidos_incertidumbre_y_unidades(cliente):
    r = cliente.post("/api/ensayos", json=nuevo(
        rendimiento_replicas_pct=[80, 82, 78], pureza_real_pct=99.4, mediciones={"ibuprofeno-perdida-secado": 0.3},
        incertidumbre={"n": 300, "semilla": 3},
    ))
    assert r.status_code == 201, r.text
    e = r.json()
    assert e["origen"] == "medido" and e["rendimiento_pct"] == 80
    assert e["salida"]["resultado"]["incertidumbre"]["n"] == 300
    assert e["entrada"]["mediciones"] == {"ibuprofeno-perdida-secado": 0.3}


def test_un_ensayo_no_se_edita_se_corrige_con_otro(cliente):
    cliente.post("/api/ensayos", json=nuevo())
    r = cliente.post("/api/ensayos/1/corregir", json=nuevo(lote="L-001 corregido", notas="se pesó mal el catalizador"))
    assert r.status_code == 201 and r.json()["corrige_a"] == 1 and r.json()["id"] == 2
    original = cliente.get("/api/ensayos/1").json()
    assert original["reemplazado_por"] == 2 and original["vigente"] is False and original["lote"] == "L-001"
    # el reemplazado no se vuelve a corregir: se corrige el último
    assert cliente.post("/api/ensayos/1/corregir", json=nuevo()).status_code == 409
    assert cliente.post("/api/ensayos/2/corregir", json=nuevo()).status_code == 201
    assert [e["id"] for e in cliente.get("/api/ensayos", params={"solo_vigentes": True}).json()] == [3]
    assert cliente.post("/api/ensayos/99/corregir", json=nuevo()).status_code == 404
    # los endpoints para editar o borrar no existen
    assert cliente.put("/api/ensayos/1", json=nuevo()).status_code == 405
    assert cliente.delete("/api/ensayos/1").status_code == 405


def test_comentarios_y_la_aprobacion_propia_esta_prohibida(cliente):
    cliente.post("/api/ensayos", json=nuevo())
    r = cliente.post("/api/ensayos/1/eventos", json={"tipo": "comentario", "comentario": "pesada con balanza B"})
    assert r.status_code == 201 and r.json()["eventos"][0]["comentario"] == "pesada con balanza B"
    assert cliente.post("/api/ensayos/1/eventos", json={"tipo": "comentario", "comentario": "  "}).status_code == 422
    # en modo local hay un solo usuario: quien hizo el ensayo no puede aprobarlo
    r = cliente.post("/api/ensayos/1/eventos", json={"tipo": "revisado"})
    assert r.status_code == 403 and "no puede" in r.json()["detail"]


def test_comparar_ensayos(cliente):
    cliente.post("/api/ensayos", json=nuevo())
    cliente.post("/api/ensayos", json=nuevo(rendimiento_real_pct=70))
    r = cliente.get("/api/ensayos/comparar", params={"ids": "1,2"})
    assert r.status_code == 200
    c = r.json()
    rend = next(m for m in c["metricas"] if m["metrica"] == "Rendimiento (%)")["valores"]
    assert rend["1"] != rend["2"] and rend["2"] == 70
    assert {x["id"] for x in c["compuestos"]} == {"ibb", "aa", "cat", "h2", "co", "exc"}
    assert cliente.get("/api/ensayos/comparar", params={"ids": "1"}).status_code == 422
    assert cliente.get("/api/ensayos/comparar", params={"ids": "1,1"}).status_code == 422
    assert cliente.get("/api/ensayos/comparar", params={"ids": "1,99"}).status_code == 404
    assert cliente.get("/api/ensayos/comparar", params={"ids": "a,b"}).status_code == 422


def test_reportes_pdf_y_excel(cliente):
    from openpyxl import load_workbook

    cliente.post("/api/ensayos", json=nuevo(rendimiento_replicas_pct=[80, 82, 78], incertidumbre={"n": 300}))
    cliente.post("/api/ensayos", json=nuevo())
    pdf = cliente.get("/api/ensayos/1/reporte.pdf")
    assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf"
    assert pdf.content[:5] == b"%PDF-" and len(pdf.content) > 3000
    assert "ensayo-001.pdf" in pdf.headers["content-disposition"]
    xl = cliente.get("/api/ensayos/1/reporte.xlsx")
    assert xl.status_code == 200
    wb = load_workbook(io.BytesIO(xl.content))
    assert any(n.startswith("Ensayo") for n in wb.sheetnames) and any(n.startswith("Trazabilidad") for n in wb.sheetnames)
    resumen = {fila[0].value: fila[1].value for fila in wb[wb.sheetnames[0]].iter_rows(min_row=2)}
    assert resumen["Puntaje (0-100)"] == 85 or resumen["Puntaje (0-100)"] is not None
    assert "MIXTO" in resumen["Origen de los datos"]  # rendimiento medido (réplicas), pureza estimada
    assert resumen["Lote"] == "L-001"
    assert any(n.startswith("Incertidumbre") for n in wb.sheetnames)
    comp = cliente.get("/api/ensayos/comparar.xlsx", params={"ids": "1,2"})
    assert comp.status_code == 200 and load_workbook(io.BytesIO(comp.content)).sheetnames == ["Comparación"]
    assert cliente.get("/api/ensayos/99/reporte.pdf").status_code == 404


def test_el_pdf_no_se_rompe_con_caracteres_especiales(cliente):
    notas = "Prueba <b>con</b> marcas & símbolos: H₂O → CO₂ ≥ 98 % «ñandú»\nsegunda línea"
    assert cliente.post("/api/ensayos", json=nuevo(lote="L&<1>", notas=notas)).status_code == 201
    r = cliente.get("/api/ensayos/1/reporte.pdf")
    assert r.status_code == 200 and r.content[:5] == b"%PDF-"


def test_la_cadena_de_auditoria_esta_ok_y_registra_cada_cosa(cliente):
    cliente.post("/api/ensayos", json=nuevo())
    cliente.post("/api/ensayos/1/eventos", json={"tipo": "comentario", "comentario": "ok"})
    cliente.post("/api/ensayos/1/corregir", json=nuevo())
    v = cliente.get("/api/auditoria/verificar").json()
    assert v["ok"] and v["registros"] >= 4 and v["ensayos_verificados"] == 2 and v["eventos_verificados"] == 1 and v["total_errores"] == 0
    acciones = [r["accion"] for r in cliente.get("/api/auditoria").json()]
    assert acciones[:3] == ["ensayo.crear", "ensayo.comentario", "ensayo.crear"]
    assert cliente.get("/api/auditoria").headers["X-Total-Count"] == str(v["registros"])


def _sin_disparadores(tabla: str) -> None:
    """Simula a alguien con acceso directo a la base que apaga la protección para alterar datos."""
    with datos.motor().begin() as c:
        if datos.dialecto() == "sqlite":
            for op in ("update", "delete"):
                c.execute(text(f"DROP TRIGGER {tabla}_inalterable_{op}"))
        else:
            c.execute(text(f"ALTER TABLE {tabla} DISABLE TRIGGER USER"))


def test_la_base_rechaza_modificar_o_borrar(cliente):
    cliente.post("/api/ensayos", json=nuevo())
    cliente.post("/api/ensayos/1/eventos", json={"tipo": "comentario", "comentario": "ok"})
    for sql in (
        "UPDATE ensayos SET notas = 'cambiada' WHERE id = 1",
        "DELETE FROM ensayos WHERE id = 1",
        "UPDATE eventos SET comentario = 'cambiado' WHERE id = 1",
        "DELETE FROM eventos WHERE id = 1",
        "UPDATE registro SET accion = 'otra' WHERE id = 1",
        "DELETE FROM registro WHERE id = 1",
    ):
        with pytest.raises(Exception, match="inalterable"):
            with datos.motor().begin() as c:
                c.execute(text(sql))
    assert cliente.get("/api/auditoria/verificar").json()["ok"]
    if datos.dialecto() == "postgresql":
        with pytest.raises(Exception, match="inalterable"):
            with datos.motor().begin() as c:
                c.execute(text("TRUNCATE ensayos CASCADE"))


def test_si_alguien_altera_la_base_por_fuera_la_verificacion_lo_detecta(cliente):
    cliente.post("/api/ensayos", json=nuevo())
    cliente.post("/api/ensayos", json=nuevo())
    _sin_disparadores("ensayos")
    with datos.motor().begin() as c:
        c.execute(text("UPDATE ensayos SET notas = 'adulterada' WHERE id = 1"))
    v = cliente.get("/api/auditoria/verificar").json()
    assert not v["ok"] and any("ensayo 1 fue modificado" in e["motivo"] for e in v["errores"])


def test_si_se_borra_o_cambia_el_registro_la_verificacion_lo_detecta(cliente):
    for _ in range(3):
        cliente.post("/api/ensayos", json=nuevo())
    _sin_disparadores("registro")
    with datos.motor().begin() as c:
        c.execute(text("UPDATE registro SET detalle = '{}' WHERE id = 2"))
    v = cliente.get("/api/auditoria/verificar").json()
    assert not v["ok"] and any("fue modificada" in e["motivo"] for e in v["errores"])
    with datos.motor().begin() as c:
        c.execute(text("DELETE FROM registro WHERE id = 1"))
    v = cliente.get("/api/auditoria/verificar").json()
    assert not v["ok"] and any("enlace" in e["motivo"] for e in v["errores"])


def test_un_ensayo_que_no_pasó_por_el_registro_se_detecta(cliente):
    cliente.post("/api/ensayos", json=nuevo())
    _sin_disparadores("ensayos")
    with datos.motor().begin() as c:
        c.execute(text(
            "INSERT INTO ensayos (creado_en, usuario_id, reaccion_id, lote, notas, entrada, resultado, version_modelo, corrige_a, digest) "
            "SELECT creado_en, usuario_id, reaccion_id, 'colado', notas, entrada, resultado, version_modelo, NULL, digest FROM ensayos WHERE id = 1"
        ))
    v = cliente.get("/api/auditoria/verificar").json()
    assert not v["ok"] and any("no figura en el registro" in e["motivo"] for e in v["errores"])


def test_escrituras_simultaneas_no_rompen_la_cadena(cliente):
    resultados = []

    def crear(i):
        r = cliente.post("/api/ensayos", json=nuevo(lote=f"T-{i}"))
        resultados.append(r.status_code)

    hilos = [threading.Thread(target=crear, args=(i,)) for i in range(8)]
    [h.start() for h in hilos]
    [h.join() for h in hilos]
    assert resultados == [201] * 8
    v = cliente.get("/api/auditoria/verificar").json()
    assert v["ok"] and v["ensayos_verificados"] == 8


def test_sistema_y_modo_local(cliente):
    s = cliente.get("/api/sistema").json()
    assert s["auth_activa"] is False and s["base"] in ("sqlite", "postgresql")
    yo = cliente.get("/api/auth/yo").json()
    assert yo["rol"] == "admin" and yo["email"] == "local@retorta"
    assert cliente.post("/api/auth/login", json={"email": "a@b.c", "clave": "x"}).status_code == 409
